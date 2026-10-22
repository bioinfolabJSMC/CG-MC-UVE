"""Build the author manuscript and submission documents from computed results.
Native Word equations use latex2mathml and the MML2OMML stylesheet installed
with Microsoft Office. Set MML2OMML_XSL to an alternative licensed installation.
Numerical results are always read from results/, never entered as invented data.
"""
from pathlib import Path
import os,re,json,copy,zipfile
from datetime import datetime,timezone
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from latex2mathml.converter import convert

ROOT=Path(__file__).resolve().parents[1]; RES=ROOT/'results';FIG=ROOT/'figures'
OUT=ROOT
TEMPLATE=ROOT/'code/resources/template.docx'
TITLE='Conditional gain regularization of Monte Carlo uninformative variable elimination for near infrared spectral calibration'
REFS=json.loads((ROOT/'references/references.json').read_text(encoding='utf8'))
REFMAP={};MATH=etree.XSLT(etree.parse(os.environ.get('MML2OMML_XSL',r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL')))
NATIVE_NS='http://schemas.openxmlformats.org/officeDocument/2006/math'

def clean_template():
    d=Document(TEMPLATE)
    for e in list(d.element.body):
        if e.tag!=qn('w:sectPr'):d.element.body.remove(e)
    # Old citation-library data and media are editable template slots.
    for rid,rel in list(d.part.rels.items()):
        if any(x in rel.reltype for x in ['/customXml','/image','/hyperlink']): del d.part.rels[rid]
    for old in list(d.settings.element.findall(qn('w:docVars'))):d.settings.element.remove(old)
    for old in list(d.settings.element.findall(qn('w:attachedTemplate'))):d.settings.element.remove(old)
    for rel in d.part.rels.values():
        if rel.reltype.endswith('/settings'):
            for rid,sr in list(rel.target_part.rels.items()):
                if sr.reltype.endswith('/attachedTemplate'):del rel.target_part.rels[rid]
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=TITLE
    d.core_properties.subject='Near infrared spectral calibration and variable selection'
    d.core_properties.created=datetime.now(timezone.utc);d.core_properties.modified=datetime.now(timezone.utc);d.core_properties.revision=1
    for old in list(d.core_properties._element.findall(qn('cp:lastPrinted'))):d.core_properties._element.remove(old)
    return d

def runs(p,text,refs=True):
    for piece in re.split(r'(\[\d+(?:[ ,\-–]+\d+)*\])',text):
        if refs and re.fullmatch(r'\[\d+(?:[ ,\-–]+\d+)*\]',piece or ''):
            old=[]
            for token in re.split('[, ]+',piece[1:-1]):
                if not token:continue
                if '-' in token or '–' in token:
                    a,b=map(int,re.split('[-–]',token));old.extend(range(a,b+1))
                else:old.append(int(token))
            for n in old:
                if n not in REFMAP:REFMAP[n]=len(REFMAP)+1
            r=p.add_run('['+','.join(str(REFMAP[n]) for n in old)+']');r.font.superscript=True
        else:p.add_run(piece)
    return p

def para(d,text,style='Paragraph',refs=True):
    p=d.add_paragraph(style=style);runs(p,text,refs)
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.widow_control=True
    if style=='Paragraph':
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(0)
    return p

def head(d,text,level=1):
    p=d.add_paragraph(text,style=f'Heading {level}')
    p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Pt(0)
    return p

def eq(d,latex,num):
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_together=True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.80),WD_TAB_ALIGNMENT.RIGHT)
    out=MATH(etree.fromstring(convert(latex).encode('utf8'))).getroot()
    # Office's MathML stylesheet can treat U+2015 as a baseline glyph.
    # Represent means with a native top bar, preserving editable mathematics.
    mq=lambda name:'{'+NATIVE_NS+'}'+name
    for acc in list(out.iter(mq('acc'))):
        char=acc.find(mq('accPr')+'/'+mq('chr'))
        if char is not None and char.get(mq('val')) in {'―','¯','‾'}:
            bar=etree.Element(mq('bar'))
            prop=etree.SubElement(bar,mq('barPr'))
            etree.SubElement(prop,mq('pos')).set(mq('val'),'top')
            bar.append(copy.deepcopy(acc.find(mq('e'))))
            acc.getparent().replace(acc,bar)
    # Wide hats are accents, not upper limits containing a literal caret.
    for lim in list(out.iter(mq('limUpp'))):
        upper=lim.find(mq('lim'))
        if upper is not None and ''.join(upper.itertext())=='^':
            acc=etree.Element(mq('acc'))
            prop=etree.SubElement(acc,mq('accPr'))
            etree.SubElement(prop,mq('chr')).set(mq('val'),'\u0302')
            acc.append(copy.deepcopy(lim.find(mq('e'))))
            lim.getparent().replace(lim,acc)
    for mr in out.iter('{'+NATIVE_NS+'}r'):
        mt=mr.find('{'+NATIVE_NS+'}t')
        if mt is not None and mt.text in {'median','sd','sign','max','MSE','RMSECV','RMSEP','RPD'}:
            rp=mr.find('{'+NATIVE_NS+'}rPr')
            if rp is None:rp=etree.Element('{'+NATIVE_NS+'}rPr');mr.insert(0,rp)
            sty=rp.find('{'+NATIVE_NS+'}sty')
            if sty is None:sty=etree.SubElement(rp,'{'+NATIVE_NS+'}sty')
            sty.set('{'+NATIVE_NS+'}val','p')
    if out.tag.endswith('oMathPara'):
        for child in list(out):
            if child.tag.endswith('oMath'):p._p.append(copy.deepcopy(child))
    else:p._p.append(out)
    p.add_run('\t('+str(num)+')')
    return p

def caption(d,text):
    p=para(d,text,style='Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(9)
    return p

def figure(d,name,text,width=5.85):
    path=FIG/(name+'.png')
    if not path.exists():raise FileNotFoundError(path)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width));caption(d,text)

def table(d,title,headers,rows,widths=None):
    c=caption(d,title);c.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [5.85/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):cells[j].text=str(v)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for key in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+key);e.set(qn('w:val'),'single' if key in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for side in ['left','right']:
                mm=OxmlElement('w:'+side);mm.set(qn('w:w'),'40');mm.set(qn('w:type'),'dxa');margins.append(mm)
            c._tc.get_or_add_tcPr().append(margins)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.5
                if len(rows)<=12:p.paragraph_format.keep_with_next=(ri<len(t.rows)-1)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
            if ri==0:
                bb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');bb.append(b);c._tc.get_or_add_tcPr().append(bb)
    return t

def grouped_table(d,title,headers,rows,widths):
    t=table(d,title,headers,rows,widths)
    start=0
    while start<len(rows):
        end=start+1
        while end<len(rows) and rows[end][0]==rows[start][0]:end+=1
        cell=t.cell(start+1,0).merge(t.cell(end,0))
        cell.text=rows[start][0];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.5
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
        start=end
    # Merging resets the anchor paragraph. Address actual XML paragraphs so
    # continuation cells do not alias the anchor and break page continuity.
    if len(rows)<=12:
        xml_rows=t._tbl.findall(qn('w:tr'))
        for ri,tr in enumerate(xml_rows):
            for pnode in tr.findall('.//'+qn('w:p')):
                prop=pnode.find(qn('w:pPr'))
                if prop is None:
                    prop=OxmlElement('w:pPr');pnode.insert(0,prop)
                keep=prop.find(qn('w:keepNext'))
                if keep is None:keep=OxmlElement('w:keepNext');prop.append(keep)
                keep.set(qn('w:val'),'1' if ri<len(xml_rows)-1 else '0')
    return t

def title_page(d):
    para(d,TITLE,'Article title')
    p=d.add_paragraph(style='Author names')
    authors=[('Cheng Guo','a'),('Yilin Zheng','b'),('Hongle An','c,*'),('Shuqi Liu','a'),('Liu Yang','a'),('Ziao Wang','a'),('Baolin Zhu','b,*'),('Cuiyan Han','a,*')]
    for i,(n,a) in enumerate(authors):
        if i:p.add_run(', ')
        p.add_run(n);p.add_run(a).font.superscript=True
    para(d,'a School of Pharmacy, Jiangsu Medical College, Jiangsu 224005, PR China','Affiliation')
    para(d,'b College of Chemistry, Nankai University, Tianjin 300071, PR China','Affiliation')
    para(d,'c Key Laboratory of Magnetic Molecules and Magnetic Information Materials (Ministry of Education), College of Chemistry and Chemical Engineering, Shanxi Normal University, Taiyuan 030006, PR China','Affiliation')
    para(d,'* Corresponding authors: Hongle An (honglean@sxnu.edu.cn); Baolin Zhu (zhubaolin@nankai.edu.cn); Cuiyan Han (hancuiyan@jsmc.edu.cn).','Correspondence details')

def introduction(d):
    head(d,'1 Introduction')
    para(d,'Near-infrared (NIR) calibration relates broad, overlapping spectral responses to a reference property through a multivariate model. Partial least-squares (PLS) regression is useful when the number of spectral variables exceeds the number of samples and neighboring wavelengths are strongly correlated.[1,2] Model performance nevertheless depends on the spectral representation. Multiplicative scatter correction (MSC), standard normal variate (SNV) transformation, and derivative filtering address different sources of variation, and their suitability should be assessed within the calibration procedure.[3–6] A preprocessing method that reduces visible scatter need not minimize prediction error for a particular analyte.')
    para(d,'Variable selection provides a second means of controlling the calibration model. Uninformative variable elimination (UVE) compares the stability of PLS regression coefficients with that of artificial noise variables.[7] Monte Carlo UVE (MC-UVE) replaces the leave-one-out coefficient calculation with repeated random calibration subsets and ranks the original variables directly.[8] This construction reduces dependence on a single fitted coefficient vector. However, its variablewise ranking does not update conditional contributions after each selection: two nearly interchangeable wavelengths can both receive large scores even when the second contributes little after the first has been selected.')
    para(d,'This limitation has motivated several complementary selection strategies. Ensemble MC-UVE aggregates random models, whereas UVE combined with the successive projections algorithm separates screening from redundancy reduction.[9–11] Competitive adaptive reweighted sampling, variable combination population analysis, iterative retention of informative variables, and bootstrapping soft shrinkage explore different variable subsets or update their sampling weights.[12–15] Reviews of PLS and NIR variable selection emphasize the diversity of these approaches and the dependence of their performance on the calibration problem.[16–18] Variable combinations, rather than isolated importance values, can be central to inverse calibration.[19,20] More recent ensemble and modified shrinkage procedures also address selection consistency and overfitting.[21,22] Interval and moving-window PLS offer an alternative when contiguous spectral regions are the preferred units of selection.[23,24]')
    para(d,'The present study develops conditional-gain MC-UVE (CG-MC-UVE), in which robust coefficient reliability determines a variable-specific penalty in a residualized one-variable regression. The resulting closed-form gain ranks each candidate conditional on the variables already selected. The retained columns are then used to fit a PLS model. The analysis follows two stages: the spectral representation is first selected by the lowest calibration RMSECV, and variable selection is subsequently performed on that representation. Data for corn oil and the active pharmaceutical ingredient (API) in tablets are used to examine prediction error, retained dimension, and stochastic variability under fixed target-ordered partitions.')

def methods(d):
    head(d,'2 Materials and methods')
    head(d,'2.1 Data sets and fixed partition',2)
    para(d,'The corn data set contains 80 samples measured on three NIR instruments and four reference properties. The spectra recorded on the m5 instrument and the oil column were used. Each spectrum contains 700 channels from 1100 to 2498 nm at 2 nm intervals. The data were distributed by Eigenvector Research (https://eigenvector.com/resources/data-sets/). The pharmaceutical data set contains the original calibration, validation, and test arrays of the 2002 International Diffuse Reflectance Conference shootout. The arrays from instrument 1 were concatenated in that order, giving 655 spectra and their API assay values. The original arrays contained 155, 40, and 460 samples. The first 597 channels, corresponding to 600–1792 nm, were retained; the remaining 53 channels were excluded as specified for this benchmark.')
    para(d,'For each data set, samples were sorted in ascending order of the target value using a stable sort. Samples at ranks 4, 8, 12, and so forth formed the prediction set; all other samples formed the calibration set. Ties retained their original source order. This gave 60 calibration and 20 prediction samples for corn and 492 calibration and 163 prediction samples for tablets. The three terminal tablet ranks remained in calibration. No samples were removed. Sample identifiers, original array membership, target rank, and final partition are supplied with the calculation files.')
    head(d,'2.2 Spectral preprocessing and PLS regression',2)
    para(d,'Raw spectra, MSC, SNV, a continuous wavelet transform (CWT), and a first derivative were compared. The MSC reference was the mean spectrum of the relevant training fold; it was never estimated using its validation or prediction samples. SNV used the mean and sample standard deviation within each spectrum. The derivative was calculated with a 15-point, second-order Savitzky–Golay filter, derivative order one, a 2 nm spacing, and polynomial interpolation at the edges. CWT used the first derivative of a Gaussian wavelet (gaus1) at scale 8, with 64 symmetrically padded channels at each boundary before the padding was removed. These fixed settings define the comparisons rather than exhaustive optimization over all possible preprocessing parameters.')
    para(d,'Each representation was first evaluated by full-spectrum PLS using five outer calibration folds and four inner folds for latent-variable selection. The preprocessing method with the lowest pooled outer RMSECV in Table 2 was selected separately for each data set. This choice was then fixed for PLS, UVE, MC-UVE, and CG-MC-UVE. Thus, every variable selector operated on the selected spectral representation. Raw spectra were included as the no-treatment candidate; selecting this candidate meant that no additional spectral preprocessing was applied before model centering.')
    para(d,'All PLS models were mean-centered with training-fold means and were fitted without variance scaling. For a centered matrix X of n samples and p variables and a centered response vector y, the PLS coefficient vector can be expressed using the weight matrix W, loading matrix P, and response loadings c.[1,2,25] The prediction uses the corresponding training means:')
    eq(d,r'\widehat{\boldsymbol\beta}_A=\mathbf W_A(\mathbf P_A^{T}\mathbf W_A)^{-1}\mathbf c_A,\qquad \widehat y=\overline y+(\mathbf x-\overline{\mathbf x})^{T}\widehat{\boldsymbol\beta}_A',1)
    para(d,'Here A is the number of latent variables. Candidate values ranged from 1 to 15, subject to the number of retained variables and the available sample rank. A PLS1 NIPALS implementation generated the coefficient path. Its predictions were checked against an independent scikit-learn PLSRegression implementation with scaling disabled.')
    head(d,'2.3 UVE and MC-UVE benchmarks',2)
    para(d,'The UVE benchmark appended p independent uniform noise columns with amplitude 10⁻¹¹ to the training spectra. Leave-one-out PLS fits with five latent variables provided a coefficient distribution for each original and noise variable. The absolute coefficient stability and noise-based retention rule were')
    eq(d,r's_j=\frac{|\overline\beta_j|}{\operatorname{sd}(\beta_{1j},\ldots,\beta_{Bj})},\qquad \mathcal S=\{j:s_j>\kappa\max_{\ell\in\mathcal N}s_\ell\}',2)
    para(d,'where N denotes the noise columns. The multiplier κ was selected from 0.3, 0.5, 0.7, 0.9, and 1.0 by calibration cross-validation. This formulation follows the noise-based reliability principle of UVE.[7,8] The implementation records the selected cutoff and any fallback for an empty retained set.')
    para(d,'For MC-UVE, B = 100 subsets were drawn without replacement, each containing one half of the available training samples, rounded down. A five-latent-variable PLS model was fitted to each subset. The absolute mean-to-sample-standard-deviation ratio in Eq. (2) ranked the original variables, and the top k were retained. The resampling count and sampling fraction reproduce the corresponding settings in the 2008 formulation.[8] The five-component selector model was fixed before the comparisons; the final PLS dimension was tuned separately. This avoids presenting the present benchmark as a numerical reproduction of the original tobacco experiments.')
    head(d,'2.4 Conditional gain regularization',2)
    para(d,'In this derivation, X contains the selected spectral representation: first-derivative corn channels or CWT tablet channels. Let βbj be the coefficient of variable j in Monte Carlo submodel b. The proposed procedure first replaces the mean and standard deviation with a median and a scaled median absolute deviation, and discounts frequent sign changes:')
    eq(d,r'm_j=\operatorname{median}_{b}(\beta_{bj}),\qquad d_j=1.4826\operatorname{median}_{b}|\beta_{bj}-m_j|',3)
    eq(d,r'h_j=\left|B^{-1}\sum_{b=1}^{B}\operatorname{sign}(\beta_{bj})\right|,\qquad q_j=\frac{|m_j|h_j}{d_j+\epsilon_\beta}',4)
    para(d,'The factor 1.4826 supplies the conventional Gaussian scale normalization; it does not assume that the resampled coefficients are independent or Gaussian. The numerical floor εβ is machine precision multiplied by the larger of the maximum coefficient standard deviation and 10⁻³⁰. A bounded weight is then formed from the median positive reliability q0:')
    eq(d,r'q_0=\operatorname{median}\{q_j:q_j>0\},\qquad w_j=\frac{q_j}{q_j+q_0+\epsilon_\beta}',5)
    para(d,'If all reliabilities are zero, q0 is set to one. The weights summarize resampling behavior and are not posterior inclusion probabilities. In particular, they do not inherit the false-selection guarantees of formal stability selection.[29,30]')
    para(d,'Suppose the current selected columns span an orthonormal basis QS. With centered xj and y, define the candidate residual rj and response residual eS as')
    eq(d,r'\mathbf P_S=\mathbf Q_S\mathbf Q_S^{T},\qquad \mathbf r_j=(\mathbf I-\mathbf P_S)\mathbf x_j,\qquad \mathbf e_S=(\mathbf I-\mathbf P_S)\mathbf y',6)
    para(d,'For the empty set, PS is zero. The penalty is proportional to the original centered column energy and inversely related to reliability:')
    eq(d,r'\lambda_j=\tau\|\mathbf x_j\|_2^2\frac{1-w_j}{\max(w_j,10^{-12})},\qquad \tau\in\{0.01,0.1,1\}',7)
    para(d,'The column-energy factor places the penalty on the scale of the residual sum of squares in the candidate direction. For each unselected variable, consider the one-dimensional penalized objective')
    eq(d,r'L_j(\alpha)=\|\mathbf e_S-\alpha\mathbf r_j\|_2^2+\lambda_j\alpha^2',8)
    para(d,'Expanding Eq. (8) and differentiating gives')
    eq(d,r'L_j(\alpha)=\mathbf e_S^{T}\mathbf e_S-2\alpha\mathbf r_j^{T}\mathbf e_S+\alpha^2(\mathbf r_j^{T}\mathbf r_j+\lambda_j)',9)
    eq(d,r'\frac{\partial L_j}{\partial\alpha}=-2\mathbf r_j^{T}\mathbf e_S+2\alpha(\mathbf r_j^{T}\mathbf r_j+\lambda_j)=0',10)
    eq(d,r'\alpha_j^*=\frac{\mathbf r_j^{T}\mathbf e_S}{\mathbf r_j^{T}\mathbf r_j+\lambda_j},\qquad G_j=L_j(0)-L_j(\alpha_j^*)=\frac{(\mathbf r_j^{T}\mathbf e_S)^2}{\mathbf r_j^{T}\mathbf r_j+\lambda_j}',11)
    para(d,'The candidate with the largest Gj is added to the selected set. The score therefore balances the response covariance remaining after projection, residual predictor energy, and Monte Carlo coefficient reliability. A weakly reliable variable receives a larger penalty. A variable lying in the selected span has rj = 0 and contributes no conditional information. Orthogonalization and stability-based screening have precedents in UVE–SPA, orthogonal forward selection, and decorrelation methods.[10,31,32] The proposed modification is the explicit use of robust MC-UVE reliability in this conditional ridge gain, not the use of orthogonalization or repeated sampling alone.')
    para(d,'After selection, the normalized residual direction uj is used to update the projection:')
    eq(d,r'\mathbf u_j=\frac{\mathbf r_j}{\|\mathbf r_j\|_2},\qquad \mathbf e_{S\cup j}=\mathbf e_S-\mathbf u_j(\mathbf u_j^{T}\mathbf e_S)',12)
    para(d,'All candidate columns are updated by the same projection, with a second projection against the new direction to limit numerical drift. Directions below the recorded rank tolerance are excluded. This update removes the full least-squares projection; it does not use the shrunk coefficient α*j. Consequently, Gj is a local ranking surrogate. The algorithm is not claimed to minimize a single global ridge objective, and a larger gain does not guarantee a lower PLS prediction error. The final PLS model is refitted to the retained columns of the chosen representation, rather than to the orthogonalized residuals. These inputs are first-derivative channels for corn and CWT channels for tablets. Supplementary Section S1 gives the boundary cases and the distinction between penalized gain and actual squared-error reduction.')
    head(d,'2.5 Model selection and performance assessment',2)
    para(d,'After preprocessing selection, five outer calibration folds evaluated each variable-selection procedure under the fixed representation. Within each outer training set, four inner folds selected the subset parameter and final PLS dimension. Learned transformation quantities, centering means, and variable rankings were estimated within the relevant training folds. The final subset and PLS dimension were selected again by four-fold cross-validation of the complete calibration set, followed by fitting and prediction-set evaluation. The spectral representation was held common across all competing selectors within each data set.')
    para(d,'This nesting separates variable selection and PLS tuning from their associated validation samples.[33–37] However, preprocessing was selected beforehand using the same calibration samples. The RMSECV values in Table 3 are therefore conditional on the selected preprocessing and do not estimate the entire two-stage procedure without selection bias. The fixed prediction set was excluded from both stages of selection and provided the subsequent assessment of prediction error. Neither RMSEP nor the retrospective prediction-error paths were used to choose a configuration.')
    para(d,'For a candidate θ, let Ef(θ) be the mean squared validation error in fold f. The pooled mean squared error identified θmin. The one-standard-error rule admitted candidates whose pooled error did not exceed the minimum plus the standard error of the fold errors at θmin:')
    eq(d,r'\operatorname{MSE}(\theta)\leq\operatorname{MSE}(\theta_{\min})+\frac{\operatorname{sd}\{E_f(\theta_{\min})\}}{\sqrt{F}}',13)
    para(d,'Among admissible candidates, the smallest mean number of selected variables was preferred, followed by the fewest latent variables and then the lowest error. The same rule was applied to the competing selectors. Fold errors are dependent, so this standard error is a model-selection heuristic rather than a formal confidence bound. The cardinality grid and all numerical settings are given in Supplementary Table S1. A centered training matrix has rank at most n − 1. CG paths used a conservative computational cap of n − 2 and stopped earlier if the residual rank tolerance was reached.')
    eq(d,r'\operatorname{RMSECV}=\sqrt{\frac{1}{n_c}\sum_{i=1}^{n_c}(y_i-\widehat y_i^{\mathrm{outer}})^2},\qquad \operatorname{RMSEP}=\sqrt{\frac{1}{n_p}\sum_{i=1}^{n_p}(y_i-\widehat y_i)^2}',14)
    eq(d,r'\operatorname{RPD}=\frac{s_{y,p}}{\operatorname{RMSEP}},\qquad R_p^2=1-\frac{\sum_i(y_i-\widehat y_i)^2}{\sum_i(y_i-\overline y_p)^2}',15)
    para(d,'In Eq. (14), nc and np are the numbers of calibration and prediction samples, respectively; the outer superscript identifies a prediction made when that calibration sample was held out. The main comparisons report RMSECV and RMSEP together with the retained dimension and number of PLS latent variables. The additional diagnostics in Eq. (15) were retained in the supplementary results and calculation files. The prediction-set standard deviation used the sample count minus one in its denominator. RPD varies inversely with RMSEP on a fixed prediction set and therefore does not provide independent evidence of prediction accuracy.[38] Prediction bias, mean absolute error, and residual standard deviation are also provided in the calculation files. Randomized selection methods can yield different subsets with similar prediction errors.[39,40] MC-UVE and CG-MC-UVE were therefore run 30 times with the calibration-selected hyperparameters held fixed. Means, sample standard deviations, and selection frequencies describe algorithmic variability. The conditional four-fold RMSECV from these runs is distinct from the second-stage RMSECV used in the main model comparison.')
    para(d,'Prediction-error paths versus feature count were generated only after the model choices were fixed and were not used for tuning. Paired percentile bootstrap intervals for differences in RMSEP used 5000 resamples of prediction sample indices. These intervals are conditional on the fitted models and this split; they do not capture a new training population or instrument. Ablations replaced the conditional criterion with robust-reliability ranking alone, or removed its reliability penalty to obtain unpenalized conditional gain.')

def reference_list(d):
    head(d,'References')
    formatted=(ROOT/'references/references_formatted.txt').read_text(encoding='utf8').splitlines()
    lookup={int(re.match(r'\[(\d+)\]',s).group(1)):re.sub(r'^\[\d+\]\s*','',s) for s in formatted if re.match(r'\[\d+\]',s)}
    for old,new in sorted(REFMAP.items(),key=lambda x:x[1]):
        p=para(d,f'[{new}] '+lookup[old],'EndNote Bibliography',refs=False)
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.left_indent=Inches(.35);p.paragraph_format.first_line_indent=Inches(-.35)
        p.paragraph_format.keep_together=True
    assert set(REFMAP)==set(range(1,42))-{26,27,28}
    maps=[dict(citation_number=new,record_id=old,title=REFS[old-1]['title'],doi=REFS[old-1].get('doi','')) for old,new in sorted(REFMAP.items(),key=lambda x:x[1])]
    pd.DataFrame(maps).to_csv(ROOT/'references/manuscript_citation_map.csv',index=False)
    (ROOT/'references/manuscript_references_formatted.txt').write_text('\n\n'.join(f'[{new}] '+lookup[old] for old,new in sorted(REFMAP.items(),key=lambda x:x[1])),encoding='utf-8')

def save(d,name):
    d.save(OUT/name)
    # No stale source ADDIN fields or bibliographic custom XML may remain.
    with zipfile.ZipFile(OUT/name) as z:
        xml=z.read('word/document.xml')
        assert b'EN.CITE' not in xml and b'EN.REFLIST' not in xml
        assert not any(s.startswith('customXml/') for s in z.namelist())

def rowof(df,ds,method):return df[(df.dataset==ds)&(df.method==method)].iloc[0]
def f(v,n=4):return f'{float(v):.{n}f}'
def model_table(d,df):
    methods=['PLS','UVE','MC-UVE','CG-MC-UVE']
    rows=[]
    for ds,label in [('corn','Corn'),('tablet','Tablets')]:
        for m in methods:
            r=rowof(df,ds,m);rows.append([label,m,int(r.n_variables),int(r.LV),f(r.RMSECV),f(r.RMSEP)])
    grouped_table(d,'Table 3. Comparison of PLS calibration models after selection of spectral preprocessing. RMSECV and RMSEP are expressed in % for corn oil and mg for tablet API.',['Data set','Method','k','LV','RMSECV','RMSEP'],rows,[.72,1.77,.45,.45,1.23,1.23])

def abstract(df):
    c=rowof(df,'corn','CG-MC-UVE');cm=rowof(df,'corn','MC-UVE');t=rowof(df,'tablet','CG-MC-UVE');tm=rowof(df,'tablet','MC-UVE')
    return (f'Monte Carlo uninformative variable elimination (MC-UVE) ranks spectral variables by coefficient stability without updating their contributions after selection. We developed conditional-gain MC-UVE (CG-MC-UVE), which combines robust coefficient reliability with a variable-specific ridge penalty in forward selection. A closed-form gain measures the contribution of each candidate after accounting for previously selected variables. Corn spectra recorded on m5 for oil prediction and tablet spectra for active pharmaceutical ingredient (API) prediction were evaluated using fixed target-ordered calibration/prediction sets of 60/20 and 492/163 samples. Minimum full-spectrum partial least-squares (PLS) RMSECV selected first-derivative corn spectra and continuous-wavelet-transformed tablet spectra before variable selection. Subsequent RMSECV estimates are conditional on these choices. Compared with MC-UVE, CG-MC-UVE reduced retained variables from {int(cm.n_variables)} to {int(c.n_variables)} and latent variables from {int(cm.LV)} to {int(c.LV)} for corn; the corresponding reductions for tablets were from {int(tm.n_variables)} to {int(t.n_variables)} and from {int(tm.LV)} to {int(t.LV)}. Tablet prediction errors were similar in magnitude, with RMSEP decreasing from {tm.RMSEP:.4f} to {t.RMSEP:.4f} mg, whereas corn RMSEP increased from {cm.RMSEP:.5f} to {c.RMSEP:.5f}%; RMSECV changed in the same directions. Across 30 runs at fixed configurations, CG-MC-UVE yielded smaller standard deviations of RMSEP and conditional RMSECV for both data sets. Subset agreement, however, decreased for corn and increased for tablets. Conditional gain offers an alternative strategy for near-infrared (NIR) variable selection, yielding more compact PLS models and less variation in error across Monte Carlo runs, with data-dependent prediction accuracy.')

def results(d,df):
    head(d,'3 Results and discussion')
    head(d,'3.1 Data audit and spectral variation',2)
    stats=pd.read_csv(RES/'dataset_summary.csv')
    para(d,'The two supplied spectral matrices contained no missing or nonfinite entries, constant columns, or exactly duplicated spectral rows. The original tablet source groups contributed 114/41, 31/9, and 347/113 calibration/prediction samples, respectively, after repartitioning. Table 1 summarizes the target distributions. Oil is expressed as percentage and API assay as milligrams, following the conventional units reported for these benchmarks.[41] Figure 1 shows the measured spectral profiles; Supplementary Figure S1 shows the target ranks and distributions after splitting.')
    rows=[]
    for _,r in stats.iterrows():rows.append([('Corn' if r.dataset=='corn' else 'Tablet')+' '+r.partition.lower(),int(r.n),int(r.p),f'{r.minimum:.3f}–{r.maximum:.3f}',f(r['mean'],3),f(r.sd,3)])
    table(d,'Table 1. Sample numbers, spectral variables, and target distributions. Corn values are in %; tablet values are in mg. SD is the sample standard deviation.',['Data set / subset','n','p','Range','Mean','SD'],rows,[1.64,.35,.4,1.32,1.02,1.12])
    para(d,'Corn calibration values ranged from 3.088 to 3.822%, and prediction values ranged from 3.176 to 3.832%. Thus, one prediction sample was slightly above the calibration maximum. It was retained to preserve the specified splitting rule. Tablet prediction values were contained within the calibration range. These ordered partitions mainly assess interpolation over the sampled concentration ranges and are not independent collections from another production batch or instrument. The raw spectra display broad, correlated profiles, making retained dimension a useful companion to prediction error rather than a substitute for it.')
    figure(d,'Figure_1_raw_spectra','Figure 1. NIR spectra from the supplied data: (A) 80 corn samples measured on m5; (B) 655 tablet samples measured on instrument 1 over the retained spectral region. Blue curves show individual spectra; black curves show their means.')
    head(d,'3.2 Comparison of spectral preprocessing',2)
    pp=pd.read_csv(RES/'preprocessing_comparison.csv')
    cpp=pp[(pp.dataset=='corn')&(pp.preprocessing=='D1')].iloc[0];tpp=pp[(pp.dataset=='tablet')&(pp.preprocessing=='CWT')].iloc[0]
    para(d,f'Table 2 identifies the first derivative as the selected preprocessing for the m5 corn spectra and CWT as the selected preprocessing for tablets. Their RMSECV values were {cpp.RMSECV:.4f}% and {tpp.RMSECV:.4f} mg, respectively, the lowest among the five candidates for each data set. All subsequent selectors therefore received first-derivative corn spectra or CWT-transformed tablet spectra. This choice was fixed before variable selection and was not revised using RMSEP. The comparison concerns the specified preprocessing settings; it does not establish the best filter or wavelet parameters across all possible settings.')
    rows=[]
    for key,label in [('corn','Corn'),('tablet','Tablets')]:
        sub=pp[pp.dataset==key]
        rows.extend([[label,{'D1':'First derivative'}.get(r.preprocessing,r.preprocessing),int(r.LV),f(r.RMSECV),f(r.RMSEP)] for _,r in sub.iterrows()])
    grouped_table(d,'Table 2. Full-spectrum PLS models for the preprocessing comparison. The lowest RMSECV determines the representation used for variable selection. RMSECV and RMSEP are expressed in % for corn oil and mg for tablet API.',['Data set','Preprocessing','LV','RMSECV','RMSEP'],rows,[.75,1.65,.45,1.5,1.5])
    para(d,'The comparison was organized sequentially. First, the full-spectrum PLS models in Table 2 determined the spectral representation. Second, the selected representation was supplied to every variable-selection method in Table 3. Five outer folds and four inner folds evaluated and tuned the second stage while retaining the same preprocessing choice. The full-spectrum PLS row in Table 3 consequently provides the corresponding selected row from Table 2 as a common baseline. Figure 2 shows the spectra entering this stage and the calibration-based preprocessing comparison.')
    figure(d,'Figure_2_preprocessing_and_PLS','Figure 2. Calibration spectra supplied to variable selection for (A) corn, using first-derivative spectra, and (B) tablets, using CWT. (C,D) Full-spectrum PLS RMSECV for the preprocessing candidates in corn and tablets. The minimum in each data set determines the fixed representation used in subsequent comparisons.')
    head(d,'3.3 Variable selection from the selected spectra',2)
    c=rowof(df,'corn','CG-MC-UVE');cm=rowof(df,'corn','MC-UVE');t=rowof(df,'tablet','CG-MC-UVE');tm=rowof(df,'tablet','MC-UVE')
    para(d,f'Table 3 compares models fitted to the representations selected in Table 2. For m5 corn spectra after first differentiation, CG-MC-UVE retained {int(c.n_variables)} channels, compared with {int(cm.n_variables)} for MC-UVE, a reduction of {100*(1-c.n_variables/cm.n_variables):.1f}%. The selected PLS dimension also decreased from {int(cm.LV)} to {int(c.LV)} latent variables. This more compact model gave a higher RMSECV ({c.RMSECV:.5f} versus {cm.RMSECV:.5f}%) and a higher RMSEP ({c.RMSEP:.5f} versus {cm.RMSEP:.5f}%). MC-UVE gave the lowest corn errors among the four models in Table 3. Fewer selected channels and latent variables therefore described a reduction in model dimension, but did not imply improved corn prediction accuracy.')
    para(d,f'For CWT-preprocessed tablets, CG-MC-UVE retained {int(t.n_variables)} channels and {int(t.LV)} latent variables, compared with {int(tm.n_variables)} channels and {int(tm.LV)} latent variables for MC-UVE. RMSECV decreased from {tm.RMSECV:.4f} to {t.RMSECV:.4f} mg and RMSEP from {tm.RMSEP:.4f} to {t.RMSEP:.4f} mg. The prediction errors were similar in magnitude, with the smaller conditional-gain model giving a {100*(1-t.RMSEP/tm.RMSEP):.1f}% lower RMSEP. This descriptive comparison does not establish statistical equivalence; the paired bootstrap interval for the RMSEP difference included zero (Supplementary Table S5). Together, the two data sets show that the dimensional reduction was consistent, while its effect on prediction accuracy depended on the calibration problem.')
    model_table(d,df)
    para(d,'In Table 3, k denotes the number of retained spectral variables; LV denotes the final PLS dimension. RMSECV evaluates variable selection and PLS tuning conditional on the preprocessing chosen in Table 2. RMSEP evaluates the final model on the fixed prediction set. The one-standard-error rule favors fewer inputs among configurations within the inner-error band; it is distinct from the minimum-RMSECV rule used to select preprocessing. For both first-derivative and CWT spectra, a selected position represents a transformed channel. Its calculation uses neighboring absorbance measurements, so the retained count does not specify the number of isolated wavelengths required for acquisition.')
    figure(d,'Figure_3_prediction_error_paths','Figure 3. Prediction RMSE versus the number of selected spectral variables for (A) corn and (B) tablets. Vertical dotted lines mark the calibration-selected counts. For CG-MC-UVE, the displayed path uses the selected penalty parameter. Each path point uses the PLS dimension with minimum calibration CV error at that count. The one-standard-error final model can use a different dimension even at the marked count. These retrospective prediction curves were not used for final model selection.')
    head(d,'3.4 Model complexity and variability',2)
    rs=pd.read_csv(RES/'repeat_summary.csv');st=pd.read_csv(RES/'selection_stability.csv')
    rc=rowof(rs,'tablet','CG-MC-UVE');rm=rowof(rs,'tablet','MC-UVE')
    rcc=rowof(rs,'corn','CG-MC-UVE');rcm=rowof(rs,'corn','MC-UVE')
    para(d,'Reducing both the retained input count and the PLS dimension produced a more compact calibration model. To assess its sensitivity to the Monte Carlo ensemble, each final configuration was repeated 30 times with new ensemble seeds. Table 4 reports the resulting errors. The preprocessing, subset size, penalty, and PLS dimension were fixed; the coefficient ensembles and variable rankings were recalculated within each training fold. These repetitions measure algorithmic variation on the same samples.')
    rows=[]
    for ds,label in [('corn','Corn'),('tablet','Tablets')]:
        for m in ['MC-UVE','CG-MC-UVE']:
            r=rowof(rs,ds,m);model=rowof(df,ds,m)
            rows.append([label,m,int(model.n_variables),int(model.LV),f'{r.conditional_RMSECV_mean:.5f} ± {r.conditional_RMSECV_sd:.5f}',f'{r.RMSEP_mean:.5f} ± {r.RMSEP_sd:.5f}'])
    grouped_table(d,'Table 4. Model dimensions and error variation over 30 runs at fixed configurations. Values are mean ± sample SD, in % for corn oil and mg for tablet API. Conditional RMSECV is distinct from the nested estimates in Table 3.',['Data set','Method','k','LV','Conditional RMSECV','RMSEP'],rows,[.70,1.35,.35,.35,1.55,1.55])
    para(d,f'CG-MC-UVE gave smaller standard deviations for both conditional RMSECV and RMSEP. For corn, RMSEP SD decreased from {rcm.RMSEP_sd:.5f} to {rcc.RMSEP_sd:.5f}%, a reduction of {100*(1-rcc.RMSEP_sd/rcm.RMSEP_sd):.1f}%; for tablets, it decreased from {rm.RMSEP_sd:.4f} to {rc.RMSEP_sd:.4f} mg, a reduction of {100*(1-rc.RMSEP_sd/rm.RMSEP_sd):.1f}%. The corresponding reductions in conditional RMSECV SD were {100*(1-rcc.conditional_RMSECV_sd/rcm.conditional_RMSECV_sd):.1f}% and {100*(1-rc.conditional_RMSECV_sd/rm.conditional_RMSECV_sd):.1f}%. These results support greater consistency of the error measures across Monte Carlo runs at the selected configurations. They do not establish robustness to new instruments, production batches, or preprocessing choices.')
    para(d,f'The lower error dispersion should be distinguished from the mean error. Corn mean RMSEP was {rcc.RMSEP_mean:.5f}% for CG-MC-UVE and {rcm.RMSEP_mean:.5f}% for MC-UVE, preserving the higher-error ordering of the single-run comparison. Tablet mean RMSEP was {rc.RMSEP_mean:.4f} and {rm.RMSEP_mean:.4f} mg, respectively. Thus, more consistent errors across random runs did not necessarily mean more accurate predictions. The conditional CV errors in Table 4 also do not replace the errors obtained with nested subset and PLS tuning in Table 3.')
    figure(d,'Figure_4_selected_channels','Figure 4. Selected channel positions on the mean calibration spectrum under the chosen representation: (A) first-derivative corn spectra and (B) CWT tablet spectra. Marker shapes distinguish MC-UVE and CG-MC-UVE. (C,D) Inclusion frequencies over 30 runs with fixed hyperparameters. Positions identify transformed channels, whose calculation uses neighboring raw observations.')
    sc=rowof(st,'corn','CG-MC-UVE');sm=rowof(st,'corn','MC-UVE');tc=rowof(st,'tablet','CG-MC-UVE');tmst=rowof(st,'tablet','MC-UVE')
    para(d,f'Figure 4 addresses a different aspect of variability: the identities of the selected channels. Chance-adjusted subset agreement decreased from {sm.chance_corrected_agreement:.3f} for MC-UVE to {sc.chance_corrected_agreement:.3f} for CG-MC-UVE in corn, but increased from {tmst.chance_corrected_agreement:.3f} to {tc.chance_corrected_agreement:.3f} in tablets (Supplementary Table S9). Lower variation in error therefore did not imply uniformly more reproducible subsets. Correlated channels can share predictive information, so the reproducibility of individual channel identities and the dispersion of prediction error should be reported separately. No chemical band assignments or instrument-transfer stability were inferred from these results.')
    ab=pd.read_csv(RES/'ablation_comparison.csv')
    cga=rowof(ab,'corn','GAIN');cra=rowof(ab,'corn','ROBUST')
    para(d,f'The corn ablations further separate the combined criterion from its components. Unpenalized conditional gain retained {int(cga.n_variables)} channels and gave an RMSEP of {cga.RMSEP:.5f}%, while robust reliability ranking retained {int(cra.n_variables)} channels and gave {cra.RMSEP:.5f}%. Both errors were below the {c.RMSEP:.5f}% obtained with {int(c.n_variables)} channels by CG-MC-UVE. Because the common calibration rule selected different input dimensions, these results do not isolate the effect of the penalty at a matched subset size.')
    ga=rowof(ab,'tablet','GAIN');ra=rowof(ab,'tablet','ROBUST')
    para(d,f'The ablations isolate two components of the selection criterion (Supplementary Table S3). For tablets, unpenalized conditional gain gave an RMSEP of {ga.RMSEP:.4f} mg, robust reliability ranking alone gave {ra.RMSEP:.4f} mg, and CG-MC-UVE gave {t.RMSEP:.4f} mg. Their retained counts were {int(ga.n_variables)}, {int(ra.n_variables)}, and {int(t.n_variables)}, respectively. Error differences therefore need to be considered together with the model dimensions selected by the common rule. These comparisons characterize the observed contributions of reliability weighting and conditional selection; they do not imply that either component always improves prediction.')
    head(d,'3.5 Scope and limitations',2)
    para(d,'Several limits define the scope of these comparisons. Preprocessing was selected from the same calibration samples used for the subsequent cross-validation, so Table 3 does not provide an unbiased error estimate for the complete two-stage procedure. The fixed prediction sets remained excluded from model selection, but their sizes were only 20 samples for corn and 163 for tablets. The bootstrap intervals are conditional on these samples and fitted models. The specified preprocessing settings, five-component selector models, and finite parameter grid also restrict the comparison. A selected grid point is not a globally minimal sufficient subset. Finally, the gain is derived for a scalar residual least-squares objective, whereas the final calibration uses PLS; improvement of the local selection objective need not improve PLS prediction.')
    para(d,'The method combines established elements of resampling reliability, shrinkage, and orthogonal selection in a specific conditional criterion. The algebra establishes the local objective decrease and the rejection of redundant directions; it does not establish statistical consistency, error control, or prediction superiority. The current evidence is therefore best interpreted as a reproducible examination of a parsimonious MC-UVE modification. It supports large reductions in retained dimension, with mixed predictive consequences.')
    para(d,'Figure 5 summarizes the two-stage procedure: preprocessing is selected by the lowest full-spectrum PLS RMSECV, variables are selected from the resulting spectra, and the final PLS model is evaluated on the held-out prediction set.')
    figure(d,'Figure_5_analysis_workflow','Figure 5. Sequential calculation workflow. The lowest full-spectrum PLS RMSECV selects the spectral representation before variable selection. Subset parameters and PLS dimensions are then tuned on the selected spectra. Prediction sets remain excluded from both selection stages. The second-stage RMSECV is conditional on the chosen preprocessing.')

def conclusion(d,df):
    head(d,'4 Conclusions')
    c=rowof(df,'corn','CG-MC-UVE');cm=rowof(df,'corn','MC-UVE');t=rowof(df,'tablet','CG-MC-UVE');tm=rowof(df,'tablet','MC-UVE')
    para(d,f'CG-MC-UVE combines robust coefficient reliability with the contribution of a candidate after previously selected variables have been accounted for. Applied to the representations selected by minimum full-spectrum PLS RMSECV, it reduced both the number of spectral inputs and the PLS dimension relative to MC-UVE: from {int(cm.n_variables)} to {int(c.n_variables)} channels and {int(cm.LV)} to {int(c.LV)} latent variables for corn, and from {int(tm.n_variables)} to {int(t.n_variables)} channels and {int(tm.LV)} to {int(t.LV)} latent variables for tablets. Tablet prediction errors remained similar in magnitude, with a lower RMSEP for CG-MC-UVE, whereas both cross-validation and prediction errors were higher for corn.')
    para(d,'Across 30 runs at fixed configurations, CG-MC-UVE gave smaller standard deviations of RMSEP and conditional RMSECV for both data sets. This improvement in the consistency of error measures did not extend uniformly to variable identities: subset agreement increased for tablets but decreased for corn. Conditional gain therefore offers an alternative approach to NIR variable selection that considers coefficient reliability and the remaining contribution of each candidate together. The results support compact PLS models with less variation in error under the specified sampling conditions, while the balance between model dimension and predictive accuracy remains data dependent. Cross-validation estimates remain conditional on preprocessing selected before variable selection.')

FUNDING='This work was supported by the Scientific Research Start-up Fund of Jiangsu Medical College (20246105); the Key R&D Program of Yancheng (YCBE202445); the Natural Science Foundation in Colleges of Jiangsu Province (25KJD150006); the 2025 Jiangsu Science and Technology Vice President Program (FZ20251560); the 2025 Jiangsu Industry–University–Research Cooperation Program (BY20250828); and the Fundamental Research Program of Shanxi Province (202503021212222).'

def declarations(d,blind=False):
    if not blind:head(d,'Acknowledgements');para(d,FUNDING)
    head(d,'Data availability')
    para(d,'The corn and pharmaceutical tablet data originate from the public Eigenvector Research and IDRC 2002 benchmark collections. The exact input-file hashes, sample partitions, analysis settings, Python source code, selected features, and numerical figure data accompany this study. No new experimental measurements were collected.')
    head(d,'Declaration of generative AI use')
    para(d,'OpenAI Codex assisted with the formulation and implementation of the candidate algorithm, numerical analysis, figure preparation, and drafting of the manuscript. The supplied observations were analyzed without adding generated experimental data. The executable analysis and its numerical outputs are provided for verification.')

def manuscript(df,blind=False):
    REFMAP.clear()
    d=clean_template()
    if not blind:title_page(d);d.add_page_break()
    para(d,TITLE,'Article title');para(d,'Abstract: '+abstract(df),'Abstract');para(d,'Keywords: near-infrared spectroscopy; variable selection; Monte Carlo UVE; conditional gain; partial least-squares regression','Keywords')
    introduction(d);methods(d);results(d,df);conclusion(d,df);declarations(d,blind);reference_list(d)
    save(d,'Manuscript_blinded.docx' if blind else 'Manuscript.docx')

def supplementary(df):
    d=clean_template();para(d,'Supplementary material','Article title');para(d,TITLE)
    para(d,'This material provides the derivation of the conditional-gain criterion, calculation settings, ablation results, repeated-run summaries, and selected feature identities. The primary model comparisons use RMSECV and RMSEP. Supplementary diagnostics characterize sampling variability and conditional prediction uncertainty; they do not replace the preprocessing-conditional error estimates in the main comparison.')
    head(d,'S1 Derivation and limiting cases')
    para(d,'Write c = rᵀe and d = rᵀr for a candidate residual direction. For d + λ > 0, the second derivative of the penalized objective is 2(d + λ), so its stationary point is the unique minimum. Substitution of α* = c/(d + λ) gives the penalized gain in the main text. The reduction in the unpenalized residual sum of squares at the same shrunk coefficient is instead')
    eq(d,r'\Delta_{\mathrm{SSE}}(\alpha^*)=2\alpha^*c-(\alpha^*)^2d=\frac{c^2(d+2\lambda)}{(d+\lambda)^2}', 'S1')
    para(d,'Subtracting the penalized gain from this reduction yields λc²/(d + λ)², which is the penalty paid at the optimum. The subsequent orthogonal projection removes the entire least-squares direction, giving a third quantity:')
    eq(d,r'\Delta_{\mathrm{projection}}=\frac{c^2}{d},\qquad G=\frac{c^2}{d+\lambda}\leq\Delta_{\mathrm{SSE}}(\alpha^*)\leq\frac{c^2}{d}', 'S2')
    para(d,'These inequalities hold for d > 0 and λ ≥ 0. The three reductions coincide when λ = 0. As λ tends to infinity, both the penalized gain and the shrunk-coefficient improvement tend to zero. If a column belongs to the selected span, r = 0 and c = 0; it is excluded by the numerical rank check. The algorithm uses G to order candidates and full projection to update the selected span. Thus, summing successive G values does not give the decrease of a global ridge objective.')
    para(d,'The penalty can also be interpreted through the angle between the residual predictor and residual response. Define ρ = c/(√d ‖e‖) for nonzero residual vectors. Then')
    eq(d,r'G=\|\mathbf e\|_2^2\rho^2\frac{d}{d+\lambda}', 'S3')
    para(d,'The gain is the remaining response energy multiplied by a squared conditional correlation and a shrinkage factor between zero and one. A large variablewise reliability is insufficient when the residual direction has little relation to the remaining response. Conversely, a strongly predictive residual direction may be discounted by an unstable Monte Carlo coefficient distribution. This decomposition describes the ranking criterion, not a probabilistic calibration of its components.')
    para(d,'For fixed reliability w, rescaling a candidate x to ax rescales c to ac, d to a²d, and λ to a²λ, leaving the gain unchanged. This is a conditional algebraic property. Re-estimating PLS coefficient reliability after rescaling the whole input matrix need not preserve w, particularly because PLS is fitted without variance scaling; the complete procedure is therefore not claimed to be invariant to arbitrary spectral rescaling.')
    para(d,'For B coefficient vectors and a maximum path length K, reliability calculation requires O(Bp) storage or equivalent streamed summaries. Given the coefficient ensemble, direct candidate projection and gain evaluation require approximately O(npK) arithmetic. PLS submodel fits and nested parameter searches contribute additional costs. These expressions are not a measured speed comparison with MC-UVE.')
    head(d,'S2 Preprocessing definitions')
    para(d,'For spectrum xi, SNV subtracts its across-channel mean and divides by its across-channel sample standard deviation. The implementation floors the divisor at 10⁻¹⁵ to avoid division by zero:')
    eq(d,r'z_{ij}=\frac{x_{ij}-\overline x_i}{\max(s_i,10^{-15})},\qquad s_i^2=\frac{1}{p-1}\sum_{j=1}^{p}(x_{ij}-\overline x_i)^2', 'S4')
    para(d,'For MSC, let μ denote the mean spectrum of the current training set and μc its centered form. Least-squares slope and intercept estimates for each spectrum are')
    eq(d,r'b_i=\frac{(\mathbf x_i-\overline x_i\mathbf 1)^T\boldsymbol\mu_c}{\boldsymbol\mu_c^T\boldsymbol\mu_c},\qquad a_i=\overline x_i-b_i\overline\mu,\qquad \mathbf z_i=\frac{\mathbf x_i-a_i\mathbf 1}{b_i}', 'S5')
    para(d,'The training reference is reused for the associated validation or prediction spectra. A slope with magnitude below 10⁻¹² raises an error; none occurred. The first derivative is the derivative at the center of a local second-order polynomial fitted to 15 adjacent channels, with wavelength expressed in nm. The edge values use polynomial interpolation. CWT is implemented by discrete convolution with the first Gaussian derivative at scale 8. It is used here as a spectral representation and is distinct from the db4 decomposition used by the WT selectors.')
    para(d,'The discrete wavelet transform (DWT) is a deterministic map W from Rᵖ to Rᑫ with periodization at each level. For corn, the A3, D3, D2, and D1 blocks contain 88, 88, 175, and 350 coefficients; for tablets they contain 75, 75, 150, and 299 coefficients. Although db4 belongs to an orthogonal wavelet family, the implemented maps for these sequence lengths are not square p × p matrices. No exact finite-dimensional energy-preservation identity is required by the selection method.')
    head(d,'S3 Validation design and calculation settings')
    settings=[['Input partition','Stable ascending target sort; each fourth rank predicts'],['Random seed','20261021'],['Outer / inner folds','5 / 4 shuffled calibration folds'],['Monte Carlo submodels','100; 50% of the current training set'],['Selector PLS dimension','5 latent variables'],['Final PLS dimension','1–15, bounded by available rank and variables'],['UVE noise / cutoff','Uniform [−1,1] × 10⁻¹¹; κ = 0.3, 0.5, 0.7, 0.9, 1'],['CG penalty parameter','τ = 0.01, 0.1, 1'],['Candidate counts','2, 4, 6, 8, 12, 16, 24, 32, 48, 64, 96, 128, 192, 256, 384, 512; feasible maximum'],['Subset selection','One-standard-error error band; fewest inputs, then fewest LV'],['Repeated runs','30 independent ensemble seeds; fixed hyperparameters'],['Bootstrap','5000 paired prediction-sample resamples'],['Software','Python; NumPy 2.3.5; SciPy 1.18.1; PyWavelets 1.8.0']]
    settings_table=table(d,'Table S1. Numerical settings.',['Setting','Specification'],settings,[1.65,4.20])
    for row in settings_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.35
    para(d,'Preprocessing selection preceded all variable selection. For each of the five representations, five outer calibration folds were evaluated, with four inner folds selecting the PLS dimension. The representation with the lowest pooled outer RMSECV was selected from main Table 2: the first derivative for corn and CWT for tablets. This choice was then fixed across the main comparisons and ablation experiments. In every subsequent outer fold, four inner folds selected the subset parameters and PLS dimension using the fixed representation. Learned preprocessing quantities, centering means, and variable rankings were recalculated from the relevant training data.')
    para(d,'Main Table 2 estimates PLS error separately for each candidate representation; its minimum determines the common input representation for main Table 3. The second-stage RMSECV nests variable selection and PLS tuning, but is conditional on a preprocessing choice made using the complete calibration set. These errors do not account for preprocessing-selection uncertainty and may be optimistic for the complete two-stage procedure. The fixed prediction samples were not used to select preprocessing, the reliability penalty, subset size, or PLS dimension. The retrospective prediction-error paths were calculated after the configurations had been fixed and were not used for additional tuning.')
    para(d,'The preprocessing choice uses the smallest RMSECV across the five candidates, without a one-standard-error band across preprocessing methods. Within each candidate, and subsequently during variable selection, the one-standard-error rule operates on mean squared error. Its threshold is the minimum pooled MSE plus the sample standard deviation of the fold MSE values at that candidate divided by the square root of the fold count. Eligible models are ordered by the smallest mean retained count, then the fewest latent variables, and then MSE. This heuristic favors compact models without assuming independent fold errors or a coverage guarantee.')
    table(d,'Table S2. Final configurations selected using the calibration set.',['Data set','Method','Prep.','k','LV','Parameter'],[[r.dataset,r.method,r.preprocessing,int(r.n_variables),int(r.LV),str(r.parameter)] for _,r in df.iterrows()],[.62,1.42,1.08,.45,.4,1.88])
    head(d,'S4 Ablation experiments')
    para(d,'ROBUST ranks variables by descending qj from the Monte Carlo coefficient ensemble without conditional forward selection. GAIN uses unpenalized conditional forward selection (τ = 0), removing the reliability-dependent penalty. Both were applied to the same first-derivative corn and CWT tablet spectra used in the main comparisons, with the same latent-variable range, cardinality grid, conditional validation, and one-standard-error rule. RPD is retained here as an ancillary diagnostic; for a fixed prediction set, it is determined by RMSEP and the response standard deviation.')
    ab=pd.read_csv(RES/'ablation_comparison.csv')
    table(d,'Table S3. Ablations of robust reliability and conditional gain.',['Data set','Method','k','LV','RMSECV','RMSEP','RPD'],[[r.dataset,r.method,int(r.n_variables),int(r.LV),f(r.RMSECV),f(r.RMSEP),f(r.RPD,3)] for _,r in ab.iterrows()],[.65,1.1,.45,.4,1.15,1.15,.95])
    para(d,'The ablations compare the unpenalized conditional path and robust reliability ranking with their combination in CG-MC-UVE. Each configuration was selected from calibration data under the same error-band rule, so the resulting dimensions can differ. Their prediction errors should be read alongside those dimensions. The comparisons use the same fixed prediction samples and do not constitute independent replications.')
    head(d,'S5 Repeated runs and conditional uncertainty')
    rs=pd.read_csv(RES/'repeat_summary.csv');assert len(rs)==8 and (rs.runs==30).all()
    for ds in ['corn','tablet']:
        s=rs[rs.dataset==ds]
        table(d,'Table S4. '+('(A) Corn' if ds=='corn' else 'Continued. (B) Tablets')+'. Thirty-run summaries; values are mean ± sample SD.',['Method','Conditional RMSECV','RMSEP','RPD'],[[r.method,f'{r.conditional_RMSECV_mean:.4f} ± {r.conditional_RMSECV_sd:.4f}',f'{r.RMSEP_mean:.4f} ± {r.RMSEP_sd:.4f}',f'{r.RPD_mean:.3f} ± {r.RPD_sd:.3f}'] for _,r in s.iterrows()],[1.55,1.55,1.50,1.25])
    para(d,'Supplementary Table S4 evaluates 30 runs at each final model configuration, with preprocessing, retained count, penalty, and PLS dimension held fixed. Main Table 4 reports the spectral MC-UVE and CG-MC-UVE subset of these same summaries. The ensemble and variable ranking are recomputed within each cross-validation training fold. These summaries quantify variation caused by submodel sampling and are not confidence intervals for population prediction accuracy. Supplementary Table S11 separately reports wavelet coefficient-count paths, with PLS dimension selected by the calibration CV minimum at each count. These conditional errors differ from the second-stage nested errors in main Table 3. Zero sample SD indicates identical numerical results over the specified runs.')
    para(d,'For M repeated subsets of the same size k, let cj be the number of subsets containing feature j and let p be the size of the relevant feature space. The average pairwise intersection and its chance-adjusted agreement are calculated as')
    eq(d,r'\overline I=\frac{\sum_{j=1}^{p}c_j(c_j-1)}{M(M-1)},\qquad A=\frac{\overline I-k^2/p}{k-k^2/p}', 'S6')
    para(d,'Identical subsets give A = 1, whereas independently and uniformly selected size-k subsets have expected A = 0. Negative values indicate less overlap than this reference. These values depend on the feature representation and do not identify whether interchangeable correlated wavelengths carry the same chemical information.')
    bs=pd.read_csv(RES/'paired_bootstrap.csv')
    table(d,'Table S5. Paired bootstrap differences in prediction RMSE relative to spectral-domain MC-UVE. Negative differences favor the listed method.',['Data set','Method','Difference','95% interval'],[[r.dataset,r.method,f(r.delta_RMSEP),f'[{r.CI_low:.4f}, {r.CI_high:.4f}]'] for _,r in bs.iterrows()],[.7,1.75,1.1,2.3])
    para(d,'Each bootstrap draw resamples the same prediction-sample indices for both methods and recalculates the difference in RMSEP. Percentiles 2.5 and 97.5 form the interval. Models are not retrained, and preprocessing remains fixed. An interval entirely above zero indicates a higher prediction error for the listed model within this conditional resampling analysis. An interval spanning zero does not establish a direction of the difference and does not demonstrate equivalence. These intervals do not include uncertainty from the calibration population or from choosing preprocessing.')
    head(d,'S6 Selected inputs and additional diagnostics')
    sel=pd.read_csv(RES/'selected_features.csv')
    rows=[]
    for ds in ['corn','tablet']:
        ss=sel[(sel.dataset==ds)&(sel.method=='CG-MC-UVE')].sort_values('selection_order')
        rows.append([ds,', '.join(str(int(x)) for x in ss.wavelength_nm)])
    table(d,'Table S6. Wavelength positions of selected CG-MC-UVE channels in selection order (nm). Positions refer to first-derivative corn channels and CWT-transformed tablet channels.',['Data set','Channel positions (nm)'],rows,[.75,5.1])
    ss=sel[sel.method=='WT-CG-MC-UVE']
    table(d,'Table S7. Selected WT-CG-MC-UVE coefficients. Global index is one-based; within-band index follows the exported metadata convention.',['Data set','Order','Global index','Band','Band index'],[[r.dataset,int(r.selection_order),int(r.index_1),r.wavelet_band,int(r.wavelet_index_in_band)] for _,r in ss.iterrows()],[.9,.65,1.4,1.3,1.6])
    para(d,'The machine-readable selected_features.csv includes every selected input for all compared models, with zero-based and one-based global indices. For corn, the listed wavelengths locate the retained first-derivative channels; for tablets, they locate the retained CWT-transformed channels. Both preprocessing operations combine neighboring raw observations. DWT indices refer to coefficients obtained after the selected preprocessing. All identities describe the final single-run model, whereas inclusion frequencies describe the separate repeated-run analysis. Correlation matrices quantify variation shared by selected inputs and do not establish chemically independent bands.')
    checks=json.loads((RES/'independent_verification.json').read_text())
    para(d,'Independent calculation checks confirmed disjoint and complete partitions, exact implementation of the fourth-rank rule, and agreement of stored PLS predictions with scikit-learn PLSRegression for all 16 spectral, wavelet, and ablation models at the recorded numerical tolerance. Additional checks verified the one-dimensional penalized-gain identity, removal of duplicate directions by the rank guard, and use of training-only MSC references. Input-file hashes and complete check results are included in independent_verification.json and implementation_checks.json.')
    table(d,'Table S8. Calculation files supporting reproducibility.',['File group','Contents'],[['*_split.csv','Sample identifiers, source groups, target ranks, partitions, outer and inner folds'],['*_model.json','Final choices, tuning paths, and outer-fold configurations'],['*_final.npz','Predictions, outer predictions, selected indices, feature matrices, and scores'],['*_curve.csv','Locked calibration and retrospective prediction paths'],['*_repeat_curves.csv','Every repeated-run conditional error and count'],['*_repeats.npz','Repeated predictions and inclusion frequencies'],['selected_features.csv','All selected wavelength and wavelet identities'],['*_wavelet_metadata.csv','Transform bands and coefficient indices'],['analysis_config.json','Seeds, grids, preprocessing, and software versions']],[1.8,4.05])
    st=pd.read_csv(RES/'selection_stability.csv')
    table(d,'Table S9. Selection agreement over 30 runs at fixed subset size.',['Data set','Method','k','Mean intersection','Agreement'],[[r.dataset,r.method,int(round(r.k)),f(r.mean_pairwise_intersection,3),f(r.chance_corrected_agreement,3)] for _,r in st.iterrows()],[.7,1.6,.45,1.6,1.5])
    head(d,'S7 Supplementary wavelet coefficient comparisons')
    p=para(d,'Discrete wavelet coefficients provide a multiresolution representation for regression and variable selection.',refs=False);p.add_run('[S1,S2]').font.superscript=True
    para(d,'The selected spectral preprocessing was followed by a level-3 db4 transform with periodization boundary handling. Approximation coefficients A3 were concatenated with detail coefficients D3, D2, and D1. The coefficient counts were 701 for corn and 599 for tablets because the implementation extends odd-length sequences. Both WT-MC-UVE and WT-CG-MC-UVE used this transform. A retained coefficient represents a combination of channels and is not a single acquisition wavelength.')
    p=para(d,'No threshold-based wavelet denoising was applied; coefficient selection and soft-threshold denoising are distinct operations.',refs=False);p.add_run('[S3]').font.superscript=True
    wt=df[df.method.str.startswith('WT-')]
    grouped_table(d,'Table S10. PLS models using selection of discrete wavelet coefficients after the fixed preprocessing. Error units are % for corn oil and mg for tablet API.',['Data set','Method','k','LV','RMSECV','RMSEP'],[['Corn' if r.dataset=='corn' else 'Tablets',r.method,int(r.n_variables),int(r.LV),f(r.RMSECV),f(r.RMSEP)] for _,r in wt.iterrows()],[.72,1.77,.45,.45,1.23,1.23])
    cw=rowof(df,'corn','WT-CG-MC-UVE');cwm=rowof(df,'corn','WT-MC-UVE');tw=rowof(df,'tablet','WT-CG-MC-UVE');twm=rowof(df,'tablet','WT-MC-UVE')
    para(d,f'The WT models applied db4 decomposition after the selected preprocessing: first differentiation for corn and CWT for tablets. For corn, WT-CG-MC-UVE retained {int(cw.n_variables)} coefficients, compared with {int(cwm.n_variables)} for WT-MC-UVE. Its RMSECV was {cw.RMSECV:.5f}%, compared with {cwm.RMSECV:.5f}%, and its RMSEP was {cw.RMSEP:.5f}%, compared with {cwm.RMSEP:.5f}%. Both errors increased despite the smaller coefficient set. Furthermore, both wavelet models had higher prediction errors than MC-UVE applied directly to the first-derivative channels. The additional wavelet representation therefore did not improve corn prediction under these settings.')
    para(d,f'For tablets, WT-CG-MC-UVE retained {int(tw.n_variables)} coefficients and gave an RMSECV of {tw.RMSECV:.4f} mg and an RMSEP of {tw.RMSEP:.4f} mg. WT-MC-UVE retained {int(twm.n_variables)} coefficients, with corresponding errors of {twm.RMSECV:.4f} and {twm.RMSEP:.4f} mg. Thus, the conditional-gain model had a lower cross-validation error but a higher prediction error than its matched wavelet control. These comparisons evaluate coefficient selection after the same CWT and db4 transformations. The coefficient count describes the number of PLS inputs. Since both transforms combine spectral channels, the smaller regression input set does not establish a reduction in the number of wavelengths required for measurement.')
    ws=pd.read_csv(RES/'wavelet_count_summary.csv')
    rows=[]
    for key,label in [('corn','Corn'),('tablet','Tablets')]:
        s=ws[ws.dataset==key];ks=sorted(set([2,4,8,16,32,int(s.n_variables.max())]));s=s[s.n_variables.isin(ks)]
        rows.extend([[label,int(r.n_variables),f(r.mean_RMSECV),f(r.SD_RMSECV),f(r.mean_RMSEP),f(r.SD_RMSEP)] for _,r in s.iterrows()])
    grouped_table(d,'Table S11. WT-CG-MC-UVE errors at representative coefficient counts over 30 runs. SD is the sample standard deviation. Error values are in % for corn oil and mg for tablet API.',['Data set','Coefficients','RMSECV mean','SD','RMSEP mean','SD'],rows,[.72,.96,1.2,.82,1.2,.95])
    para(d,'Figure S7 and Table S11 show how the conditional cross-validation and prediction errors varied with the number of retained wavelet coefficients. At each count, the penalty and the PLS dimension selected by calibration tuning were held fixed across 30 runs, while the variable rankings were recalculated. The PLS dimension for these paths minimizes the calibration CV error at the specified count and may differ from the final one-standard-error choice. Accordingly, the conditional RMSECV values in Table S11 should be distinguished from the preprocessing-conditional estimates in Supplementary Table S10. Complete paths, run-level errors, and coefficient identities are provided in the Supplementary calculation files.')
    rs=pd.read_csv(RES/'repeat_summary.csv');rwc=rowof(rs,'tablet','WT-CG-MC-UVE');rwm=rowof(rs,'tablet','WT-MC-UVE')
    para(d,f'For the tablet wavelet models, mean RMSEP over 30 runs was {rwc.RMSEP_mean:.4f} mg for WT-CG-MC-UVE and {rwm.RMSEP_mean:.4f} mg for WT-MC-UVE. This ordering differed from the single-run result in Table S10. The repetitions characterize Monte Carlo sampling variation at fixed settings on the same samples; they do not evaluate new sample populations.')
    head(d,'Supplementary references')
    formatted=(ROOT/'references/references_formatted.txt').read_text(encoding='utf-8').splitlines()
    lookup={int(re.match(r'\[(\d+)\]',line).group(1)):re.sub(r'^\[\d+\]\s*','',line) for line in formatted if re.match(r'\[\d+\]',line)}
    supmap=[];suprefs=[]
    for j,old in enumerate([26,27,28],1):
        txt=f'[S{j}] '+lookup[old];suprefs.append(txt)
        p=para(d,txt,'EndNote Bibliography',refs=False);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
        supmap.append(dict(citation_number=f'S{j}',record_id=old,title=REFS[old-1]['title'],doi=REFS[old-1].get('doi','')))
    pd.DataFrame(supmap).to_csv(ROOT/'references/supplementary_citation_map.csv',index=False)
    (ROOT/'references/supplementary_references_formatted.txt').write_text('\n\n'.join(suprefs),encoding='utf-8')
    for name,cap in [
      ('Figure_S1_partition_audit','Figure S1. Target-ordered partitions and target distributions for corn and tablets. Every fourth sorted observation is assigned to prediction.'),
      ('Figure_S2_all_preprocessing','Figure S2. Calibration spectra for all preprocessing candidates. Each row corresponds to a representation; columns correspond to corn and tablets. Learned MSC references use calibration samples only.'),
      ('Figure_S3_calibration_error_paths','Figure S3. Four-fold calibration tuning error versus selected spectral-variable count. These are four-fold tuning paths, distinct from the five-fold errors reported in the main comparison.'),
      ('Figure_S4_selected_feature_correlations','Figure S4. Pearson correlation matrices of the final selected spectral channels. Rows show corn and tablets; columns show MC-UVE and CG-MC-UVE.'),
      ('Figure_S5_wavelet_features','Figure S5. Mean transformed calibration coefficients and final WT-CG-MC-UVE selections for (A) corn and (B) tablets. (C,D) Selection frequencies over 30 runs at fixed hyperparameters. Indices refer to the concatenated DWT coefficient vector.'),
      ('Figure_S6_stochastic_variation','Figure S6. (A,B) Conditional RMSECV and (C,D) prediction RMSE over 30 runs at fixed hyperparameters for corn and tablets, respectively. MC, CG, WT-MC, and WT-CG denote MC-UVE, CG-MC-UVE, WT-MC-UVE, and WT-CG-MC-UVE. The distributions describe algorithmic sampling variability.'),
      ('Figure_S7_wavelet_repeated_CV','Figure S7. Mean conditional RMSECV ± sample SD from 30 runs versus retained wavelet coefficient count for (A) corn and (B) tablets. The penalty and PLS dimension associated with each count were fixed by calibration tuning. Variable rankings were recalculated within each validation training fold.')]:
        d.add_page_break();figure(d,name,cap)
    save(d,'Supplementary.docx')

def cover_letter(df):
    c=rowof(df,'corn','CG-MC-UVE');cm=rowof(df,'corn','MC-UVE');t=rowof(df,'tablet','CG-MC-UVE');tm=rowof(df,'tablet','MC-UVE')
    d=clean_template();para(d,'22 October 2026',refs=False);para(d,'Dear Editor,',refs=False)
    para(d,'Please consider our manuscript, “'+TITLE+'”, for publication as an original research article in Analytical Letters.',refs=False)
    para(d,'The study derives a conditional-gain criterion for Monte Carlo uninformative variable elimination (MC-UVE). Robust coefficient reliability determines a variable-specific penalty, and residual projection measures the contribution of each candidate after other variables have been selected. The resulting scalar gain guides spectral variable selection for partial least-squares (PLS) calibration.',refs=False)
    para(d,'The lowest full-spectrum PLS cross-validation error selects first-derivative m5 corn spectra and continuous-wavelet-transformed tablet spectra. Variable selection then uses those representations, with calibration-only tuning. Downstream cross-validation is conditional on the preprocessing choice; prediction samples are excluded from both selection stages.',refs=False)
    para(d,f'Compared with MC-UVE, the conditional-gain method (CG-MC-UVE) reduces retained channels from {int(cm.n_variables)} to {int(c.n_variables)} for corn and from {int(tm.n_variables)} to {int(t.n_variables)} for tablets. PLS dimensions decrease from {int(cm.LV)} to {int(c.LV)} and from {int(tm.LV)} to {int(t.LV)}, respectively. Tablet prediction errors remain similar in magnitude, while corn errors increase. Both data sets show smaller standard deviations of prediction and conditional cross-validation error across 30 runs at fixed configurations. The manuscript distinguishes error consistency from subset reproducibility and offers a strategy for compact near-infrared calibration models.',refs=False)
    para(d,'The accompanying materials include supplementary derivations, numerical results, complete analysis and plotting programs, and the data underlying the figures. The manuscript discloses the use of generative AI in its preparation.',refs=False)
    p=para(d,'Sincerely,\nHongle An, Baolin Zhu, and Cuiyan Han\nCorresponding authors\nhonglean@sxnu.edu.cn; zhubaolin@nankai.edu.cn; hancuiyan@jsmc.edu.cn',refs=False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    save(d,'Cover_letter.docx')

def ancillary(df):
    c=rowof(df,'corn','CG-MC-UVE');cm=rowof(df,'corn','MC-UVE');t=rowof(df,'tablet','CG-MC-UVE');tm=rowof(df,'tablet','MC-UVE')
    d=clean_template();para(d,'Highlights','Article title')
    items=[
        'Minimum PLS RMSECV fixes preprocessing before spectral variable selection.',
        'Conditional gain combines coefficient reliability with residual contribution.',
        f'Selected inputs decrease from {int(cm.n_variables)} to {int(c.n_variables)} for corn and {int(tm.n_variables)} to {int(t.n_variables)} for tablets.',
        f'PLS latent variables decrease from {int(cm.LV)} to {int(c.LV)} for corn and {int(tm.LV)} to {int(t.LV)} for tablets.',
        'Error standard deviations decrease over 30 runs at fixed configurations.']
    for text in items:
        assert len(text)<=85,(len(text),text)
        para(d,'• '+text,refs=False)
    save(d,'Highlights.docx')
    d=clean_template()
    for i in range(1,4):
        if i>1:d.add_page_break()
        para(d,'Graphical abstract '+str(i),'Article title')
        figure(d,'Graphical_abstract_'+str(i),'')
    save(d,'Graphical_abstracts.docx')

if __name__=='__main__':
    df=pd.read_csv(RES/'model_comparison.csv')
    assert set(df.loc[df.dataset=='corn','preprocessing'])=={'D1'}, 'The m5 manuscript requires the selected first-derivative representation.'
    assert set(df.loc[df.dataset=='tablet','preprocessing'])=={'CWT'}, 'The tablet comparison requires the preserved CWT representation.'
    print('Abstract word count:',len(abstract(df).split()))
    manuscript(df);supplementary(df);cover_letter(df);ancillary(df)
    print('Created five revised DOCX documents with',len(REFMAP),'cited references.')
