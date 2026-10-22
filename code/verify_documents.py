"""Audit revision 05 document contents against unchanged results and figure assets.

This reads documents and saved results only. It does not fit models, regenerate
figures, modify numerical files, or replace full-page visual inspection.
"""
from pathlib import Path
import hashlib
import json
import re
import zipfile

from docx import Document
from lxml import etree
import pandas as pd
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'results'
OUT = ROOT / 'audits'
LABEL = {'corn': 'Corn', 'tablet': 'Tablets'}
F = lambda value, places=4: f'{float(value):.{places}f}'


def cells(table, include_header=False):
    return [[c.text for c in row.cells] for row in table.rows[int(not include_header):]]


def check_table(table, header, expected, name):
    assert cells(table, True)[0] == header, f'{name}: header mismatch'
    actual = cells(table)
    assert len(actual) == len(expected), f'{name}: row-count mismatch'
    for number, (row, source) in enumerate(zip(actual, expected), 1):
        assert row == list(map(str, source)), (name, number, row, source)


def paragraphs(doc):
    return '\n'.join(p.text for p in doc.paragraphs)


def citation_ids(text):
    ids = set()
    for item in re.findall(r'\[([0-9,\-–]+)\]', text):
        for part in item.split(','):
            if '-' in part or '–' in part:
                lo, hi = map(int, re.split('[-–]', part))
                ids.update(range(lo, hi + 1))
            else:
                ids.add(int(part))
    return ids


def main():
    doc = Document(ROOT / 'Manuscript.docx')
    sup = Document(ROOT / 'Supplementary.docx')
    assert len(doc.tables) == 4 and len(sup.tables) == 12
    ds = pd.read_csv(DATA / 'dataset_summary.csv')
    pp = pd.read_csv(DATA / 'preprocessing_comparison.csv')
    models = pd.read_csv(DATA / 'model_comparison.csv')
    repeats = pd.read_csv(DATA / 'repeat_summary.csv')
    rows = [[('Corn' if r.dataset == 'corn' else 'Tablet') + ' ' + r.partition.lower(), int(r.n), int(r.p),
             f'{r.minimum:.3f}–{r.maximum:.3f}', F(r['mean'], 3), F(r.sd, 3)] for _, r in ds.iterrows()]
    check_table(doc.tables[0], ['Data set / subset', 'n', 'p', 'Range', 'Mean', 'SD'], rows, 'Table 1')
    rows = [[LABEL[r.dataset], {'D1': 'First derivative'}.get(r.preprocessing, r.preprocessing),
             int(r.LV), F(r.RMSECV), F(r.RMSEP)] for _, r in pp.iterrows()]
    check_table(doc.tables[1], ['Data set', 'Preprocessing', 'LV', 'RMSECV', 'RMSEP'], rows, 'Table 2')
    nonwt = models[models.method.isin(['PLS', 'UVE', 'MC-UVE', 'CG-MC-UVE'])]
    rows = [[LABEL[r.dataset], r.method, int(r.n_variables), int(r.LV), F(r.RMSECV), F(r.RMSEP)]
            for _, r in nonwt.iterrows()]
    check_table(doc.tables[2], ['Data set', 'Method', 'k', 'LV', 'RMSECV', 'RMSEP'], rows, 'Table 3')
    assert len(rows) == 8
    rows = []
    for dataset in ['corn', 'tablet']:
        for method in ['MC-UVE', 'CG-MC-UVE']:
            r = repeats[(repeats.dataset == dataset) & (repeats.method == method)].iloc[0]
            m = models[(models.dataset == dataset) & (models.method == method)].iloc[0]
            rows.append([LABEL[dataset], method, int(m.n_variables), int(m.LV),
                         f'{r.conditional_RMSECV_mean:.5f} ± {r.conditional_RMSECV_sd:.5f}',
                         f'{r.RMSEP_mean:.5f} ± {r.RMSEP_sd:.5f}'])
    check_table(doc.tables[3], ['Data set', 'Method', 'k', 'LV', 'Conditional RMSECV', 'RMSEP'], rows, 'Table 4')
    for dataset in ['corn', 'tablet']:
        subset = pp[pp.dataset == dataset]
        selected = subset.loc[subset.RMSECV.idxmin()]
        assert (models.loc[models.dataset == dataset, 'preprocessing'] == selected.preprocessing).all()
        pls = models[(models.dataset == dataset) & (models.method == 'PLS')].iloc[0]
        assert abs(pls.RMSECV - selected.RMSECV) < 1e-10

    # Supplementary tables S2-S11, including all numerical data moved from main text.
    rows = [[r.dataset, r.method, r.preprocessing, int(r.n_variables), int(r.LV), str(r.parameter)]
            for _, r in models.iterrows()]
    check_table(sup.tables[1], ['Data set', 'Method', 'Prep.', 'k', 'LV', 'Parameter'], rows, 'Table S2')
    ablation = pd.read_csv(DATA / 'ablation_comparison.csv')
    rows = [[r.dataset, r.method, int(r.n_variables), int(r.LV), F(r.RMSECV), F(r.RMSEP), F(r.RPD, 3)]
            for _, r in ablation.iterrows()]
    check_table(sup.tables[2], ['Data set', 'Method', 'k', 'LV', 'RMSECV', 'RMSEP', 'RPD'], rows, 'Table S3')
    for table, dataset in zip(sup.tables[3:5], ['corn', 'tablet']):
        rows = [[r.method, f'{r.conditional_RMSECV_mean:.4f} ± {r.conditional_RMSECV_sd:.4f}',
                 f'{r.RMSEP_mean:.4f} ± {r.RMSEP_sd:.4f}', f'{r.RPD_mean:.3f} ± {r.RPD_sd:.3f}']
                for _, r in repeats[repeats.dataset == dataset].iterrows()]
        check_table(table, ['Method', 'Conditional RMSECV', 'RMSEP', 'RPD'], rows, f'Table S4 {dataset}')
    bootstrap = pd.read_csv(DATA / 'paired_bootstrap.csv')
    rows = [[r.dataset, r.method, F(r.delta_RMSEP), f'[{r.CI_low:.4f}, {r.CI_high:.4f}]']
            for _, r in bootstrap.iterrows()]
    check_table(sup.tables[5], ['Data set', 'Method', 'Difference', '95% interval'], rows, 'Table S5')
    selection = pd.read_csv(DATA / 'selected_features.csv')
    rows = []
    for dataset in ['corn', 'tablet']:
        selected = selection[(selection.dataset == dataset) & (selection.method == 'CG-MC-UVE')].sort_values('selection_order')
        rows.append([dataset, ', '.join(str(int(v)) for v in selected.wavelength_nm)])
    check_table(sup.tables[6], ['Data set', 'Channel positions (nm)'], rows, 'Table S6')
    rows = [[r.dataset, int(r.selection_order), int(r.index_1), r.wavelet_band, int(r.wavelet_index_in_band)]
            for _, r in selection[selection.method == 'WT-CG-MC-UVE'].iterrows()]
    check_table(sup.tables[7], ['Data set', 'Order', 'Global index', 'Band', 'Band index'], rows, 'Table S7')
    assert any(row[0] == '*_wavelet_metadata.csv' for row in cells(sup.tables[8]))
    stability = pd.read_csv(DATA / 'selection_stability.csv')
    rows = [[r.dataset, r.method, int(round(r.k)), F(r.mean_pairwise_intersection, 3), F(r.chance_corrected_agreement, 3)]
            for _, r in stability.iterrows()]
    check_table(sup.tables[9], ['Data set', 'Method', 'k', 'Mean intersection', 'Agreement'], rows, 'Table S9')
    rows = [[LABEL[r.dataset], r.method, int(r.n_variables), int(r.LV), F(r.RMSECV), F(r.RMSEP)]
            for _, r in models[models.method.str.startswith('WT-')].iterrows()]
    check_table(sup.tables[10], ['Data set', 'Method', 'k', 'LV', 'RMSECV', 'RMSEP'], rows, 'Table S10')
    assert len(rows) == 4
    wavelets = pd.read_csv(DATA / 'wavelet_count_summary.csv')
    rows = []
    for dataset in ['corn', 'tablet']:
        subset = wavelets[wavelets.dataset == dataset]
        counts = {2, 4, 8, 16, 32, int(subset.n_variables.max())}
        rows.extend([[LABEL[dataset], int(r.n_variables), F(r.mean_RMSECV), F(r.SD_RMSECV), F(r.mean_RMSEP), F(r.SD_RMSEP)]
                     for _, r in subset[subset.n_variables.isin(counts)].iterrows()])
    check_table(sup.tables[11], ['Data set', 'Coefficients', 'RMSECV mean', 'SD', 'RMSEP mean', 'SD'], rows, 'Table S11')
    assert len(rows) == 12

    text, suptext = paragraphs(doc), paragraphs(sup)
    assert re.findall(r'^Figure (\d+)\.', text, re.M) == list(map(str, range(1, 6)))
    assert set(re.findall(r'Figure (\d+)\b', text)) == set(map(str, range(1, 6)))
    assert re.findall(r'^Table (\d+)\.', text, re.M) == list(map(str, range(1, 5)))
    assert re.findall(r'^Figure S(\d+)\.', suptext, re.M) == list(map(str, range(1, 8)))
    assert re.findall(r'^\[(\d+)\]', text, re.M) == list(map(str, range(1, 39)))
    assert re.findall(r'^\[S(\d+)\]', suptext, re.M) == ['1', '2', '3']
    assert citation_ids(text.split('\nReferences\n')[0]) == set(range(1, 39))
    assert '[S1,S2]' in suptext and '[S3]' in suptext
    assert '3.4 Model complexity and variability' in text
    assert not re.search(r'^3\.4.*wavelet', text, re.M | re.I)
    assert 'WT-CG-MC-UVE' not in text and 'WT-MC-UVE' not in text
    assert 'S7 Supplementary wavelet coefficient comparisons' in suptext
    assert 'first derivative for tablets' not in text
    assert 'preprocessing selection is also repeated' not in text
    assert 'first-derivative corn spectra' in text and 'CWT tablet spectra' in text
    assert 'mp5' not in text and 'raw corn' not in text.lower()
    assert 'conditional on the selected preprocessing' in text
    for name, content, marker in [('manuscript', text, r'^\[\d+\]'), ('supplementary', suptext, r'^\[S\d+\]')]:
        expected = [line.strip() for line in (ROOT / 'references' / f'{name}_references_formatted.txt').read_text(encoding='utf-8').splitlines() if re.match(marker, line)]
        observed = [line.strip() for line in content.splitlines() if re.match(marker, line)]
        assert observed == expected, f'{name}: bibliography text mismatch'

    names = {
        'Manuscript': ['Figure_1_raw_spectra', 'Figure_2_preprocessing_and_PLS', 'Figure_3_prediction_error_paths', 'Figure_4_selected_channels', 'Figure_5_analysis_workflow'],
        'Supplementary': ['Figure_S1_partition_audit', 'Figure_S2_all_preprocessing', 'Figure_S3_calibration_error_paths', 'Figure_S4_selected_feature_correlations', 'Figure_S5_wavelet_features', 'Figure_S6_stochastic_variation', 'Figure_S7_wavelet_repeated_CV'],
        'Highlights': [], 'Cover_letter': [],
        'Graphical_abstracts': [f'Graphical_abstract_{i}' for i in range(1, 4)],
    }
    images = sorted((ROOT / 'figures').glob('*.png'))
    assert len(images) == 15
    hashes = {p.stem: hashlib.sha256(p.read_bytes()).hexdigest() for p in images}
    for path in images:
        with Image.open(path) as image:
            assert all(abs(v - 300) < .02 for v in image.info['dpi']), path.name
    report = {
        'status': 'PASS',
        'all_audited_table_values_match_saved_CSVs': True,
        'minimum_RMSECV_preprocessing': True,
        'main_tables': 4, 'main_figures': 5, 'main_references': 38,
        'main_table_3_data_rows': 8, 'main_table_4_data_rows': 4,
        'supplementary_table_10_data_rows': 4, 'supplementary_table_11_data_rows': 12,
        'supplementary_figures': 7, 'supplementary_references': 3,
        'old_main_wavelet_section_removed': True,
        'raster_figures': len(images), 'raster_DPI': 300,
        'pages': {}, 'editable_equations': {}, 'embedded_figures': {},
    }
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    for name, expected in names.items():
        with zipfile.ZipFile(ROOT / f'{name}.docx') as archive:
            media = {hashlib.sha256(archive.read(k)).hexdigest() for k in archive.namelist() if k.startswith('word/media/')}
            assert media == {hashes[item] for item in expected}, f'{name}: embedded image hash mismatch'
            xml = etree.fromstring(archive.read('word/document.xml'))
            relationships = etree.fromstring(archive.read('word/_rels/document.xml.rels'))
            targets = {r.get('Id'): r.get('Target') for r in relationships}
            ordered = [hashlib.sha256(archive.read('word/' + targets[blip.get('{' + ns['r'] + '}embed')])).hexdigest()
                       for blip in xml.findall('.//a:blip', ns)]
            assert ordered == [hashes[item] for item in expected], f'{name}: figure order mismatch'
            equations = len(xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
            report['editable_equations'][name] = equations
            report['embedded_figures'][name] = len(media)
            assert not any(k.startswith('customXml/') for k in archive.namelist())
        report['pages'][name] = len(PdfReader(ROOT / f'{name}.pdf').pages)
    assert report['editable_equations']['Manuscript'] == 15
    assert report['editable_equations']['Supplementary'] == 6
    assert report['pages']['Highlights'] == 1 and report['pages']['Cover_letter'] == 1
    assert report['pages']['Graphical_abstracts'] == 3
    highlights = [p.text[2:] for p in Document(ROOT / 'Highlights.docx').paragraphs if p.text.startswith('• ')]
    assert len(highlights) == 5 and all(len(item) <= 85 for item in highlights)
    report['highlights_within_85_characters'] = True
    preservation = json.loads((DATA / 'results_preservation_check.json').read_text(encoding='utf-8'))
    assert preservation['all_computational_results_byte_identical']
    report['preserved_computational_result_files'] = preservation['expected_file_count']
    report['note'] = 'Structural and numerical audit; full rendered-page visual QA is recorded separately.'
    OUT.mkdir(exist_ok=True)
    (OUT / 'document_audit.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
